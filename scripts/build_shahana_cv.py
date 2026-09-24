from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "documents" / "Shahana_Batool_CV.docx"

INK = "17212B"
TEAL = "0F766E"
MUTED = "52606D"
PALE = "E7F4F2"
WHITE = "FFFFFF"
FONT = "Aptos"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_layout(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def font(run, size=9.2, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def paragraph(cell, text="", *, size=9.2, color=INK, bold=False, italic=False,
              before=0, after=0, line=1.0, align=None, keep=False):
    p = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    font(p.add_run(text), size, color, bold, italic)
    return p


def heading(cell, text, before=8):
    p = paragraph(cell, text.upper(), size=8.5, color=TEAL, bold=True, before=before, after=4, keep=True)
    p.paragraph_format.keep_together = True
    return p


def bullet(cell, text, *, after=1.5):
    p = paragraph(cell, size=8.7, color=MUTED, after=after, line=1.05)
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    font(p.add_run("•  "), 8.7, TEAL, bold=True)
    font(p.add_run(text), 8.7, MUTED)
    return p


def role(cell, organization, location, title, period, bullets):
    p = paragraph(cell, before=2, after=0, keep=True)
    font(p.add_run(organization), 10.1, INK, bold=True)
    if location:
        font(p.add_run(f"  |  {location}"), 8.4, MUTED)
    p = paragraph(cell, after=2, keep=True)
    font(p.add_run(title), 9.2, TEAL, bold=True)
    if period:
        font(p.add_run(f"  ·  {period}"), 8.4, MUTED)
    for item in bullets:
        bullet(cell, item)


def add_skill_pills(cell, skills):
    table = cell.add_table(rows=0, cols=2)
    set_repeat_table_layout(table)
    table.columns[0].width = Cm(3.2)
    table.columns[1].width = Cm(3.2)
    for index in range(0, len(skills), 2):
        row = table.add_row()
        for col in range(2):
            target = row.cells[col]
            set_cell_shading(target, PALE)
            set_cell_margins(target, 75, 90, 75, 90)
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index + col < len(skills):
                p = target.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                font(p.add_run(skills[index + col]), 8.1, INK, bold=True)
            if col == 0:
                target.width = Cm(3.2)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.05)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title_style = doc.styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style_ppr = title_style._element.get_or_add_pPr()
    title_style_border = title_style_ppr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_ppr.remove(title_style_border)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(1)
    title.paragraph_format.line_spacing = 1
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(title.add_run("Shahana Batool"), 25, "000000", bold=True)
    remove_paragraph_border(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    font(subtitle.add_run("PROJECT COORDINATOR  ·  SUPPLY CHAIN AND ADMINISTRATION"), 10.5, TEAL, bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(8)
    font(contact.add_run("Rawalpindi, Pakistan  ·  +92 318 512 6272  ·  shahanabatool149@gmail.com"), 8.7, MUTED)

    columns = doc.add_table(rows=1, cols=2)
    set_repeat_table_layout(columns)
    columns.columns[0].width = Cm(11.4)
    columns.columns[1].width = Cm(6.6)
    left, right = columns.rows[0].cells
    left.width = Cm(11.4)
    right.width = Cm(6.6)
    set_cell_margins(left, 0, 0, 0, 210)
    set_cell_margins(right, 0, 210, 0, 0)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    heading(left, "Professional Experience", before=0)
    role(left, "BeHope", "Islamabad, Pakistan", "Project Coordinator", "Jun 2024 - Jun 2026", [
        "Coordinated administrative, procurement, and logistics activities for drop-shipping and in-store healthcare operations serving Canada and North America.",
        "Worked across internal teams, vendors, and customers to support timely project delivery and resolve day-to-day operational issues.",
        "Helped streamline supply chain and administrative workflows while maintaining organized communication and follow-up.",
    ])
    role(left, "Starco Logistics", "Islamabad, Pakistan", "Supply Chain Intern", "", [
        "Supported operations within an end-to-end logistics company and gained practical exposure to supply chain coordination.",
    ])
    role(left, "Pakistan Stock Exchange Company Analysis", "NUML", "Financial Analysis Project", "Feb 2023 - Mar 2023", [
        "Analyzed listed companies using financial information and business performance indicators.",
        "Completed focused analysis of Fauji Fertilizer and Engro Foods.",
    ])
    role(left, "Entrepreneurial Journey", "NUML", "Project Participant", "Oct 2022", [
        "Developed UniRide, a shuttle-service concept focused on improving access to dependable transportation.",
    ])

    heading(left, "Education", before=7)
    role(left, "National University of Modern Languages", "Islamabad", "Bachelor of Business Administration", "Feb 2021 - Jan 2025", [
        "GPA 3.96 / 4.00  ·  Specialization in Supply Chain Management",
    ])
    role(left, "FG Postgraduate College for Women", "Wah Cantt", "FSc Pre-Medical", "2019", [
        "Percentage 86%",
    ])

    heading(right, "Profile", before=0)
    paragraph(right,
        "Business Administration graduate specializing in supply chain management, with hands-on experience in project coordination, procurement, logistics, administrative support, and financial analysis. Known for organized execution, thoughtful problem-solving, and clear coordination across teams, vendors, and customers.",
        size=8.8, color=MUTED, after=4, line=1.08)

    heading(right, "Core Skills", before=5)
    add_skill_pills(right, [
        "Project coordination", "Procurement", "Inventory management", "Logistics support",
        "Administrative support", "Financial analysis", "Strategic planning", "Time management",
        "Problem-solving", "Communication",
    ])

    heading(right, "Key Achievement", before=7)
    paragraph(right, "3rd Prize  ·  SDGs Business Startup Competition", size=9.2, color=INK, bold=True, after=2, keep=True)
    paragraph(right,
        "Recognized for an oyster-mushroom cultivation concept designed around a controlled environment, supporting sustainable agriculture and food production.",
        size=8.6, color=MUTED, after=3, line=1.08)

    heading(right, "Languages", before=6)
    for language, level in (("English", "Advanced"), ("Urdu", "Native"), ("Saraiki", "Native"), ("Chinese", "Beginner")):
        p = paragraph(right, after=1.5)
        font(p.add_run(language), 8.8, INK, bold=True)
        font(p.add_run(f"  ·  {level}"), 8.5, MUTED)

    heading(right, "Strength", before=6)
    paragraph(right, "Business acumen", size=9.1, color=INK, bold=True, after=1)
    paragraph(right,
        "Combines supply chain knowledge with administrative discipline, analytical thinking, planning, and sound decision-making.",
        size=8.6, color=MUTED, line=1.08)

    doc.core_properties.title = "Shahana Batool Professional CV"
    doc.core_properties.subject = "Project coordination supply chain and administration"
    doc.core_properties.author = "Shahana Batool"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
